"""Bounded CSV/XLSX intake: review first, then atomically create scoped drafts."""

import csv
import io
import json
import zipfile
from datetime import timedelta

from fastapi import APIRouter, Request, UploadFile
from pydantic import ValidationError
from sqlalchemy import select

from app.api import schemas
from app.api.deps import CurrentPrincipal, DbSession, request_id_of
from app.api.v1.routers.fdrs import _resolve_institution
from app.core.authorization import conflict, invalid, not_found
from app.core.security import (
    deobfuscate_store,
    lookup_hash,
    mask_fdr_number,
    obfuscate_store,
)
from app.db.base import as_utc, new_id, utcnow
from app.db.locking import transaction_lock
from app.db.models import Fdr, ImportBatch
from app.domains.audit.service import record_audit
from app.domains.risk.engine import sync_findings

router = APIRouter(tags=["bulk import"])
COLUMNS = tuple(schemas.FdrDraftIn.model_fields)


def read_rows(payload: bytes, filename: str) -> list[dict]:
    if len(payload) > 2 * 1024 * 1024:
        raise invalid("Import files must be at most 2 MB.")
    if filename.lower().endswith(".csv"):
        try:
            reader = csv.DictReader(io.StringIO(payload.decode("utf-8-sig")))
            headers = reader.fieldnames or []
            rows = []
            for row in reader:
                rows.append(row)
                if len(rows) > 500:
                    raise invalid("Import at most 500 rows at a time.")
        except (UnicodeError, csv.Error) as exc:
            raise invalid("Use a UTF-8 CSV with column headers.") from exc
    elif filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        try:
            with zipfile.ZipFile(io.BytesIO(payload)) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > 20 * 1024 * 1024
                ):
                    raise invalid("The expanded workbook exceeds 20 MB.")
            workbook = load_workbook(
                io.BytesIO(payload), read_only=True, data_only=False
            )
            try:
                sheet = workbook.active
                iterator = sheet.iter_rows(values_only=True)
                headers = [str(value or "").strip() for value in next(iterator)]
                rows = []
                for values in iterator:
                    if not any(value is not None for value in values):
                        continue
                    if any(
                        isinstance(value, str) and value.startswith("=")
                        for value in values
                    ):
                        raise invalid(
                            "Replace spreadsheet formulas with their values before importing."
                        )
                    rows.append(dict(zip(headers, values)))
                    if len(rows) > 500:
                        raise invalid("Import at most 500 rows at a time.")
            finally:
                workbook.close()
        except (zipfile.BadZipFile, ValueError, StopIteration, KeyError) as exc:
            raise invalid("This workbook could not be read.") from exc
    else:
        raise invalid("Choose a .csv or .xlsx file.")
    if len(headers) != len(set(headers)) or not set(headers).issubset(COLUMNS):
        raise invalid("Use unique supported headers: " + ", ".join(COLUMNS))
    if not rows:
        raise invalid("The import has no data rows.")
    return rows


@router.get("/imports/fdrs/template")
def template(principal: CurrentPrincipal):
    from fastapi import Response

    principal.require("fdr:draft")
    return Response(
        ",".join(COLUMNS) + "\n",
        media_type="text/csv",
        headers={
            "Content-Disposition": 'attachment; filename="fdr-import-template.csv"'
        },
    )


@router.post("/imports/fdrs/preview")
async def preview(
    file: UploadFile, principal: CurrentPrincipal, db: DbSession, request: Request
):
    principal.require("fdr:draft")
    institution = _resolve_institution(db, principal)
    raw_rows = read_rows(await file.read(2 * 1024 * 1024 + 1), file.filename or "")
    rows, review, seen = [], [], set()
    for index, raw in enumerate(raw_rows, start=2):
        errors = []
        try:
            cleaned = {
                key: value for key, value in raw.items() if value not in (None, "")
            }
            payload = schemas.FdrDraftIn.model_validate(cleaned)
            token = lookup_hash(payload.issuer_bank, payload.fdr_number)
            duplicate = (
                token in seen
                or db.scalar(
                    select(Fdr.id).where(
                        Fdr.institution_id == institution.id,
                        Fdr.fdr_lookup_hash == token,
                    )
                )
                is not None
            )
            seen.add(token)
            if duplicate:
                errors.append("Duplicate candidate: review this reference separately.")
            data = payload.model_dump(mode="json")
            if not errors:
                rows.append(data)
            review.append(
                {
                    "row": index,
                    "reference": mask_fdr_number(payload.fdr_number),
                    "bank": payload.issuer_bank,
                    "principal": str(payload.principal_amount),
                    "errors": errors,
                }
            )
        except ValidationError as exc:
            errors = [
                ".".join(map(str, item["loc"])) + ": " + item["msg"]
                for item in exc.errors(include_input=False)
            ]
            review.append(
                {
                    "row": index,
                    "reference": "—",
                    "bank": "—",
                    "principal": "—",
                    "errors": errors,
                }
            )
    batch = ImportBatch(
        id=new_id(),
        actor_id=principal.user_id,
        institution_id=institution.id,
        encrypted_rows=obfuscate_store(json.dumps(rows)),
        expires_at=utcnow() + timedelta(minutes=30),
    )
    db.add(batch)
    record_audit(
        db,
        aggregate_type="import",
        aggregate_id=batch.id,
        action="import.previewed",
        actor=principal,
        new_state={"valid_rows": len(rows), "total_rows": len(review)},
        request_id=request_id_of(request),
    )
    db.commit()
    return {
        "batch_id": batch.id,
        "valid_rows": len(rows),
        "total_rows": len(review),
        "rows": review,
        "expires_at": batch.expires_at,
    }


@router.post("/imports/fdrs/{batch_id}/commit")
def commit(batch_id: str, principal: CurrentPrincipal, db: DbSession, request: Request):
    principal.require("fdr:draft")
    transaction_lock(db, "bulk-import")
    batch = db.scalar(
        select(ImportBatch)
        .where(ImportBatch.id == batch_id, ImportBatch.actor_id == principal.user_id)
        .with_for_update()
    )
    if batch is None or batch.institution_id not in principal.institution_ids:
        raise not_found()
    if batch.status == "IMPORTED":
        return json.loads(batch.result_json)
    if as_utc(batch.expires_at) < utcnow():
        raise conflict("This preview expired. Upload the file again.")
    created = []
    for raw in json.loads(deobfuscate_store(batch.encrypted_rows)):
        payload = schemas.FdrDraftIn.model_validate(raw)
        token = lookup_hash(payload.issuer_bank, payload.fdr_number)
        if db.scalar(
            select(Fdr.id).where(
                Fdr.institution_id == batch.institution_id, Fdr.fdr_lookup_hash == token
            )
        ):
            raise conflict(
                "An instrument was added since preview. Upload again to review duplicates."
            )
        data = payload.model_dump(exclude={"fdr_number"})
        fdr = Fdr(
            id=new_id(),
            institution_id=batch.institution_id,
            created_by=principal.user_id,
            fdr_lookup_hash=token,
            fdr_number_masked=mask_fdr_number(payload.fdr_number),
            fdr_number_ciphertext=obfuscate_store(payload.fdr_number),
            **data,
        )
        db.add(fdr)
        db.flush()
        sync_findings(db, fdr)
        record_audit(
            db,
            aggregate_type="fdr",
            aggregate_id=fdr.id,
            action="fdr.imported",
            actor=principal,
            new_state={"batch_id": batch.id, "status": "DRAFT"},
            request_id=request_id_of(request),
        )
        created.append(fdr.id)
    batch.status = "IMPORTED"
    result = {
        "created": len(created),
        "fdr_ids": created,
        "note": "Drafts only. Upload evidence and attest each instrument before scrutiny.",
    }
    batch.result_json = json.dumps(result)
    batch.encrypted_rows = obfuscate_store("[]")
    db.commit()
    return result
